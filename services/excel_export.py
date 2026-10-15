"""Styled, streaming XLSX export. No photos, arbitrary paths or Excel formulas."""
from datetime import date
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo
from flask import current_app
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from database import db
from database.models import Attendance, Child, Group
from services.statistics import statistics
from services.groups import group_statistics, children_report
from services.localization import t

HEADER_FILL = PatternFill('solid', fgColor='6E473B')
ROW_FILL = PatternFill('solid', fgColor='F4E6D2')
HEADER_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
BODY_FONT = Font(name='Calibri', color='291C0E', size=11)
MONEY = '#,##0.00'


def table(workbook, key, headers, widths, formats=None):
    sheet = workbook.create_sheet(t(key))
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = '6E473B'
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    sheet.print_title_rows = '1:1'
    sheet.sheet_properties.outlinePr.summaryRight = False
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = Worksheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for index,(label,width) in enumerate(zip(headers,widths),1):
        sheet.column_dimensions[get_column_letter(index)].width = min(48,max(width,min(len(label)+2,32)))
    sheet.row_dimensions[1].height = 42
    counter = [0]
    def append(values, header=False):
        counter[0] += 1
        cells = []
        for index,value in enumerate(values):
            if isinstance(value,str): value = ILLEGAL_CHARACTERS_RE.sub('',value)
            cell = WriteOnlyCell(sheet,value=value)
            # Even a name beginning '=' remains literal text, never executable.
            if isinstance(value,str): cell.data_type = 's'
            cell.font = HEADER_FONT if header else BODY_FONT
            cell.alignment = Alignment(vertical='center',wrap_text=True)
            if header: cell.fill = HEADER_FILL
            elif counter[0]%2 == 0: cell.fill = ROW_FILL
            if not header and isinstance(value,Decimal): cell.number_format = MONEY
            if not header and formats and index in formats: cell.number_format = formats[index]
            cells.append(cell)
        sheet.append(cells)
    append(headers,True)
    def finish():
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1,counter[0])}'
    return sheet, append, finish


def build_workbook(start, end, group=None):
    group_id = group.id if group else None
    summary = statistics(start,end,group_id)
    children = children_report(start,end,group_id)
    groups = group_statistics(start,end,group_id)
    name_width = max([len(child.full_name)+2 for child,_,_,_ in children]+[24])
    group_width = max([len(row['group'].name)+2 for row in groups]+[20])
    workbook = Workbook(write_only=True)
    workbook.properties.creator = 'KinderTrack'
    workbook.properties.title = t('export_summary')
    sheet, append, _ = table(workbook,'export_summary',[t('export_metric'),t('export_value')],[44,48])
    start_date,end_date = date.fromisoformat(start),date.fromisoformat(end)
    for label,value in [('export_period',f'{start_date:%d.%m.%Y} — {end_date:%d.%m.%Y}'),
                        ('group',group.name if group else t('all_groups')),
                        ('total_children',summary['total']),('export_unique',summary['selected']['unique']),
                        ('attendance_count',summary['selected']['count'])]:
        append([t(label),value])
    append([t('accrued_period'),summary['selected']['amount']])
    append([t('average_attendance'),summary['average']])
    append([t('export_note'),t('group_membership_note')])

    _, append, finish = table(workbook,'export_attendance',
        ['№',t('date'),t('time'),t('export_child_id'),t('full_name'),t('group'),t('historical_tariff'),t('export_accrued')],
        [8,14,12,14,name_width,group_width,24,24],{1:'dd.mm.yyyy',2:'hh:mm:ss',6:MONEY,7:MONEY})
    query = select(Attendance,Child,Group.name).join(Child,Attendance.child_id==Child.id).join(Group,Child.group_id==Group.id)\
        .where(Attendance.attendance_date.between(start_date,end_date)).order_by(Attendance.attendance_date,Attendance.created_at,Attendance.id)
    if group_id is not None: query=query.where(Child.group_id==group_id)
    zone=ZoneInfo(current_app.config['TIMEZONE'])
    # Streaming avoids creating hundreds of thousands of in-memory worksheet cells.
    for index,(attendance,child,group_name) in enumerate(db.session.execute(query.execution_options(yield_per=1000)),1):
        local_time=attendance.created_at.astimezone(zone).time().replace(tzinfo=None)
        append([index,attendance.attendance_date,local_time,child.id,child.full_name,group_name,attendance.amount,attendance.amount])
    finish()
    _, append, finish = table(workbook,'export_children',
        ['ID',t('full_name'),t('group'),t('period_visits'),t('accrued_period'),t('current_tariff'),t('status')],
        [12,name_width,group_width,28,26,22,18],{4:MONEY,5:MONEY})
    for child,group_name,count,amount in children:
        append([child.id,child.full_name,group_name,count,amount,child.attendance_price,t(child.status)])
    finish()
    _, append, finish = table(workbook,'export_groups',
        [t('group'),t('total_children'),t('export_present_children'),t('attendance_count'),t('accrued_period')],
        [group_width,22,30,25,26],{4:MONEY})
    for row in groups:
        append([row['group'].name,row['total'],row['unique'],row['visits'],row['accrued']])
    finish()
    output=BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
