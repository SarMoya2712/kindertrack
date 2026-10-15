"""Real PostgreSQL group membership and safe, filter-consistent Excel exports."""
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from database import db
from database.models import Group, Child, Attendance
from services.localization import CATALOGS
from conftest import add_child, csrf, child_form
from test_attendance import scan
from test_reports import contexts


@pytest.fixture
def group_data(authorized,app,fake_embedding,monkeypatch):
    moment=[datetime(2026,10,14,0,5,tzinfo=ZoneInfo('Asia/Tashkent'))]
    for module in ('statistics','groups','attendance','scanner'):
        monkeypatch.setattr('services.'+module+'.now',lambda:moment[0])
    for name in ('Группа А','Группа Б'):
        assert authorized.post('/groups/new',data={'csrf_token':csrf(authorized),'name':name}).status_code==302
    for group_id,name,price in [(2,'Первый','25000'),(2,'Второй','12000.25'),(3,'Третий','30000.50')]:
        assert add_child(authorized,price=price,overrides={'first_name':name,'group_id':str(group_id)}).status_code==302
    selected=[1]
    monkeypatch.setattr(app.extensions['face'],'compare',lambda feature:dict(child_id=selected[0],confidence=.99,status='recognized'))
    for child_id,day_offset in [(1,0),(2,-1),(3,0)]:
        selected[0]=child_id
        moment[0]=datetime(2026,10,14,0,5,tzinfo=ZoneInfo('Asia/Tashkent'))+timedelta(days=day_offset)
        assert scan(authorized).json['status']=='success'
    moment[0]=datetime(2026,10,14,0,5,tzinfo=ZoneInfo('Asia/Tashkent'))
    assert authorized.post('/children/3/status',data={'csrf_token':csrf(authorized),'status':'inactive'}).status_code==302
    with app.app_context():
        db.session.get(Child,1).attendance_price=Decimal('40000');db.session.commit()
    return moment


def render(client,app,url):
    with contexts(app) as values: response=client.get(url)
    assert response.status_code==200
    return values[-1]


def workbook(client,query):
    response=client.get('/statistics/export.xlsx',query_string=query)
    assert response.status_code==200
    assert response.mimetype=='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'attachment;' in response.headers['Content-Disposition']
    assert 'no-store' in response.headers['Cache-Control']
    return load_workbook(BytesIO(response.data)),response


def test_group_roster_transfer_and_history(authorized,app,group_data):
    by_id={row['group'].id:row for row in render(authorized,app,'/groups?period=today')['rows']}
    assert (by_id[2]['total'],by_id[2]['active'],by_id[2]['present'],by_id[2]['absent'])==(2,2,1,1)
    assert (by_id[3]['total'],by_id[3]['active'],by_id[3]['present'],by_id[3]['absent'])==(1,0,1,0)
    detail=render(authorized,app,'/groups/2?period=today')
    assert {row[0].id for row in detail['rows']}=={1,2}
    assert detail['summary']['percentage']==Decimal('50.0')
    for suffix,expected in [('status=all',{1,2}),('status=active',{1,2}),('status=inactive',set()),
                            ('status=present',{1}),('status=absent',{2}),('q=Первый',{1}),('q=2',{2}),('q=%25',set())]:
        assert {row[0].id for row in render(authorized,app,'/groups/2?period=today&'+suffix)['rows']}==expected
    with app.app_context():
        before=[(a.id,a.child_id,a.attendance_date,a.amount,a.created_at) for a in db.session.scalars(select(Attendance).order_by(Attendance.id))]
    form=dict(child_form('12000.25'),group_id='3',first_name='Второй',csrf_token=csrf(authorized))
    assert authorized.post('/children/2/edit',data=form).status_code==302
    assert {row[0].id for row in render(authorized,app,'/groups/2')['rows']}=={1}
    assert {row[0].id for row in render(authorized,app,'/groups/3')['rows']}=={2,3}
    by_id={row['group'].id:row for row in render(authorized,app,'/groups?period=7days')['rows']}
    assert by_id[2]['total']==1 and by_id[3]['total']==2
    assert by_id[2]['accrued']==Decimal('25000') and by_id[3]['accrued']==Decimal('42000.75')
    with app.app_context():
        after=[(a.id,a.child_id,a.attendance_date,a.amount,a.created_at) for a in db.session.scalars(select(Attendance).order_by(Attendance.id))]
        assert before==after


@pytest.mark.parametrize('period,start,end',[
    ('today','2026-10-14','2026-10-14'),('yesterday','2026-10-13','2026-10-13'),
    ('7days','2026-10-08','2026-10-14'),('30days','2026-09-15','2026-10-14'),
    ('month','2026-10-01','2026-10-14'),('previous_month','2026-09-01','2026-09-30'),
    ('year','2026-01-01','2026-10-14'),('custom','2026-10-13','2026-10-14')])
def test_excel_matches_all_filters_and_statistics(authorized,app,group_data,period,start,end):
    for group_id in ('','2','3'):
        query=dict(period=period,start=start,end=end,group_id=group_id)
        with contexts(app) as captured: assert authorized.get('/statistics',query_string=query).status_code==200
        stats=captured[-1]['stats']
        report,response=workbook(authorized,query)
        assert report.sheetnames==['Общая статистика','Посещения','По детям','По группам']
        summary=dict(report.worksheets[0].iter_rows(min_row=2,values_only=True))
        for label,value in [('total_children',stats['total']),('export_unique',stats['selected']['unique']),
                            ('attendance_count',stats['selected']['count']),('accrued_period',stats['selected']['amount']),
                            ('average_attendance',stats['average'])]:
            assert Decimal(str(summary[CATALOGS['ru'][label]]))==value
        rows=list(report['Посещения'].iter_rows(min_row=2,values_only=True))
        assert len(rows)==stats['selected']['count']
        assert sum((Decimal(str(row[7])) for row in rows),Decimal(0))==stats['selected']['amount']
        for row in rows:
            assert start<=row[1].date().isoformat()<=end
            assert row[2].hour==0 and row[2].minute==5
            assert row[6]==row[7]
            if group_id: assert row[5]==('Группа А' if group_id=='2' else 'Группа Б')
        children=list(report['По детям'].iter_rows(min_row=2,values_only=True))
        assert len(children)==stats['total']
        assert sum((Decimal(str(row[4])) for row in children),Decimal(0))==stats['selected']['amount']
        for row in children:
            if row[0]==1: assert row[5]==40000
            if group_id: assert row[2]==('Группа А' if group_id=='2' else 'Группа Б')
        groups=list(report['По группам'].iter_rows(min_row=2,values_only=True))
        assert sum(row[3] for row in groups)==stats['selected']['count']
        assert sum((Decimal(str(row[4])) for row in groups),Decimal(0))==stats['selected']['amount']
        if group_id: assert len(groups)==1
        assert f'{start}_{end}.xlsx' in response.headers['Content-Disposition']
        for sheet in report.worksheets:
            assert sheet.freeze_panes=='A2' and sheet['A1'].font.bold and not sheet._images
        for sheet in report.worksheets[1:]: assert sheet.auto_filter.ref
        if rows:
            assert report['Посещения']['G2'].number_format=='#,##0.00'
            assert report['Посещения']['B2'].number_format=='dd.mm.yyyy'
        report.close()


def test_excel_literal_strings_localization_access_and_no_paths(authorized,app,group_data):
    with app.app_context():
        db.session.get(Child,1).last_name='=HYPERLINK("https://example.invalid","x")'
        db.session.get(Group,2).name='=1+1';db.session.commit()
    for lang in CATALOGS:
        authorized.post('/language/'+lang,data={'csrf_token':csrf(authorized)})
        report,_=workbook(authorized,dict(period='today',group_id='2'))
        assert report.sheetnames==[CATALOGS[lang][key] for key in ('export_summary','export_attendance','export_children','export_groups')]
        assert report.worksheets[1]['E2'].value.startswith('=HYPERLINK') and report.worksheets[1]['E2'].data_type=='s'
        assert report.worksheets[1]['F2'].value=='=1+1' and report.worksheets[1]['F2'].data_type=='s'
        assert report.worksheets[0]['B3'].data_type=='s'
        assert CATALOGS[lang]['download_excel'] in authorized.get('/statistics').text
        assert CATALOGS[lang]['view_group'] in authorized.get('/groups').text
        report.close()
    assert app.test_client().get('/statistics/export.xlsx').status_code==302
    assert app.test_client().get('/groups/2').status_code==302
    for value in ('999','../../.env',"1 OR 1=1",'1'*50):
        assert authorized.get('/statistics',query_string={'group_id':value}).status_code==404
        assert authorized.get('/statistics/export.xlsx',query_string={'group_id':value}).status_code==404
    assert authorized.get('/groups/999').status_code==404


def test_empty_group_export_and_group_choices(authorized,app,group_data):
    report,_=workbook(authorized,dict(period='today',group_id='1'))
    assert report['Посещения'].max_row==1 and report['По детям'].max_row==1
    assert report['По группам']['B2'].value==0
    assert render(authorized,app,'/groups/1')['summary']['percentage']==0
    assert render(authorized,app,'/statistics?group_id=1')['stats']['selected']['amount']==0
    authorized.post('/groups/2/status',data={'csrf_token':csrf(authorized),'status':'inactive'})
    assert 'Группа А' in authorized.get('/statistics').text
    assert 'Группа А' not in authorized.get('/children/new').text
    assert 'Группа А' in authorized.get('/children/1/edit').text


def test_export_snapshot_stays_consistent_during_new_attendance(authorized,app,group_data,monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from services import excel_export
    original=excel_export.children_report
    def insert_during_export(*args,**kwargs):
        def insert_visit():
            with app.app_context():
                db.session.add(Attendance(child_id=2,attendance_date=group_data[0].date(),amount=Decimal('12000.25')))
                db.session.commit()
        with ThreadPoolExecutor(max_workers=1) as pool: pool.submit(insert_visit).result(timeout=5)
        return original(*args,**kwargs)
    monkeypatch.setattr(excel_export,'children_report',insert_during_export)
    report,_=workbook(authorized,dict(period='today'))
    summary=dict(report.worksheets[0].iter_rows(min_row=2,values_only=True))
    assert summary[CATALOGS['ru']['attendance_count']]==2
    assert report['Посещения'].max_row==3
    assert sum(row[3] for row in report['По детям'].iter_rows(min_row=2,values_only=True))==2
    assert render(authorized,app,'/statistics?period=today')['stats']['selected']['count']==3
